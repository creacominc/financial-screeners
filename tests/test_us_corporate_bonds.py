import csv
from dataclasses import replace
from datetime import date
from decimal import Decimal as D
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from gilt_analyser.calculator import _xnpv
from us_corporate_analyser.calculator import (
    CorporateBond, UnaffordableCorporateBond, accrued_interest, analyse_bond, coupon_schedule, validate,
)
from us_corporate_analyser.reporting import COLUMNS, read_bonds, ranked_rows, write_reports
from us_corporate_analyser.screening import exclusion_reason, maturity_limit, SUPPORTED_TERMS

SETTLEMENT = date(2026, 10, 19)
ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / 'examples/us_corporate_bonds_illustrative.csv'


@pytest.fixture
def bond():
    return CorporateBond(cusip='DEMO00001', name='Fictional bond', coupon_rate=D('6'),
                        maturity_date=date(2027, 10, 15), clean_buy_price=D('99'),
                        accrued_interest_per_100=D('0.06666666666666666666666666667'),
                        price_settlement_date='2026-10-19', moodys='A2', sp='A-',
                        ratings_date='2026-10-16', currency='USD', terms=SUPPORTED_TERMS,
                        quote_date='2026-10-16', price_source='Fictional test only', maximum_nominal=D('10000'))


def test_cashflows_costs_and_yield(bond):
    result = analyse_bond(bond, SETTLEMENT, D('5000'), D('10'))
    assert result.nominal_purchased == D('5000')
    assert result.clean_consideration == D('4950')
    assert result.accrued_interest_per_100 == bond.accrued_interest_per_100
    assert result.accrued_interest_paid == D('5000') * bond.accrued_interest_per_100 / 100
    assert result.dirty_price_per_100 == 99 + bond.accrued_interest_per_100
    assert result.total_initial_cost == result.clean_consideration + result.accrued_interest_paid + 10
    assert result.unused_cash == 5000 - result.total_initial_cost
    assert [f.payment_date for f in result.cashflows] == [date(2027, 4, 15), date(2027, 10, 15)]
    assert [f.coupon for f in result.cashflows] == [D('150'), D('150')]
    assert [f.redemption for f in result.cashflows] == [D('0'), D('5000')]
    assert result.coupon_income == 300
    assert result.redemption_gain == 50
    assert result.redemption_proceeds == 5000
    assert result.total_proceeds == 5300
    assert result.total_cash_profit == 5300 - result.total_initial_cost
    assert abs(_xnpv(result.annualised_yield, result.total_initial_cost, result.cashflows, SETTLEMENT)) < D('0.00000001')
    assert result.annualised_yield < analyse_bond(bond, SETTLEMENT, D('5000'), D('0')).annualised_yield


def test_coupon_date_and_schedule(bond):
    day = date(2026, 10, 15)
    adjusted = replace(bond, accrued_interest_per_100=D('0'), price_settlement_date=day.isoformat())
    assert accrued_interest(adjusted, day) == 0
    assert coupon_schedule(adjusted, day) == (date(2027, 4, 15), date(2027, 10, 15))
    with pytest.raises(ValueError, match='coupon-date'):
        accrued_interest(bond, day)
    endmonth = replace(bond, maturity_date=date(2028, 8, 31))
    assert coupon_schedule(endmonth, date(2027, 8, 31)) == (date(2028, 2, 29), date(2028, 8, 31))
    with pytest.raises(ValueError, match='before maturity'):
        coupon_schedule(bond, bond.maturity_date)


def test_affordability_grid_fee_and_offer_size(bond):
    bond = replace(bond, minimum_nominal=D('2000'), nominal_increment=D('1000'))
    result = analyse_bond(bond, SETTLEMENT, D('4950'), D('10'))
    assert result.nominal_purchased == 4000
    assert result.total_initial_cost <= result.investment_budget
    capped = analyse_bond(replace(bond, maximum_nominal=D('3500')), SETTLEMENT, D('10000'), D('0'))
    assert capped.nominal_purchased == 3000
    with pytest.raises(UnaffordableCorporateBond):
        analyse_bond(bond, SETTLEMENT, D('1000'), D('0'))
    exact = replace(bond, clean_buy_price=D('100'), accrued_interest_per_100=D('0'))
    assert analyse_bond(exact, SETTLEMENT, D('2001'), D('1')).nominal_purchased == 2000
    with pytest.raises(UnaffordableCorporateBond):
        analyse_bond(exact, SETTLEMENT, D('2000.99'), D('1'))


def test_independent_one_year_yield(bond):
    start = date(2026, 10, 15)
    security = replace(bond, accrued_interest_per_100=D('0'), price_settlement_date=str(start), clean_buy_price=D('100'))
    result = analyse_bond(security, start, D('1000'), D('0'))
    # Actual/365 differs slightly from equal half-year compounding; verify the dated equation independently.
    r = result.annualised_yield
    expected = D('30') / (1+r) ** (D('182')/365) + D('1030') / (1+r)
    assert abs(expected - 1000) < D('.00000001')
    assert D('.06') < r < D('.062')


@pytest.mark.parametrize('changes,fragment', [
    ({'sp': 'BB+'}, 'below minimum'), ({'moodys': 'Ba1'}, 'below minimum'),
    ({'sp': 'BBB- (negative)'}, 'Unknown'), ({'moodys': '', 'sp': 'NR'}, 'No recognised'),
    ({'terms': ''}, 'Unsupported'), ({'terms': 'make_whole_callable'}, 'Unsupported'),
    ({'currency': 'GBP'}, 'Unsupported'), ({'coupon_rate': D('0')}, 'Unsupported'),
    ({'quote_date': '2026-10-01'}, 'Stale'), ({'quote_date': '2026-10-20'}, 'future'),
    ({'ratings_date': '2026-06-01'}, 'Stale'), ({'ratings_date': '2026-10-20'}, 'future'),
    ({'price_settlement_date': '2026-10-20'}, 'settlement'),
    ({'maturity_date': SETTLEMENT}, 'Outside'), ({'maturity_date': date(2028, 10, 20)}, 'Outside'),
])
def test_exclusions(bond, changes, fragment):
    assert fragment in exclusion_reason(replace(bond, **changes), SETTLEMENT)


def test_rating_and_calendar_boundaries(bond):
    assert exclusion_reason(bond, SETTLEMENT) == ''
    assert exclusion_reason(replace(bond, sp=''), SETTLEMENT) == ''
    assert exclusion_reason(replace(bond, moodys='NR'), SETTLEMENT) == ''
    assert 'below' in exclusion_reason(bond, SETTLEMENT, minimum_moodys='Aa3')
    assert exclusion_reason(replace(bond, maturity_date=date(2028, 10, 19)), SETTLEMENT) == ''
    assert maturity_limit(date(2027, 1, 31), 1) == date(2027, 2, 28)
    assert exclusion_reason(replace(bond, quote_date='2026-10-12', ratings_date='2026-07-21'), SETTLEMENT) == ''
    for kwargs in [{'months': 25}, {'months': 0}, {'max_quote_age': -1}, {'max_rating_age': -1}, {'minimum_sp': 'BB+'}]:
        with pytest.raises(ValueError):
            exclusion_reason(bond, SETTLEMENT, **kwargs)


@pytest.mark.parametrize('field,value', [
    ('coupon_rate', D('NaN')), ('clean_buy_price', D('0')), ('clean_buy_price', 99.0),
    ('minimum_nominal', D('-1')), ('nominal_increment', D('0')), ('maximum_nominal', D('1')),
    ('accrued_interest_per_100', D('Infinity')), ('accrued_interest_per_100', D('-1')),
    ('broker_ytm_pct', D('NaN')), ('cusip', ''), ('quote_date', ''), ('ratings_date', 'bad'),
])
def test_invalid_inputs(bond, field, value):
    with pytest.raises(ValueError):
        validate(replace(bond, **{field: value}))


@pytest.mark.parametrize('budget,fee', [(D('0'), D('0')), (D('5000'), D('-1')), (D('5'), D('5')), (D('NaN'), D('0'))])
def test_invalid_analysis_arguments(bond, budget, fee):
    with pytest.raises(ValueError):
        analyse_bond(bond, SETTLEMENT, budget, fee)


def test_unsupported_terms_and_accrual_never_analysed(bond):
    for changed in [replace(bond, terms='callable'), replace(bond, currency='GBP'),
                    replace(bond, price_settlement_date='2026-10-20'), replace(bond, accrued_interest_per_100=D('4'))]:
        with pytest.raises(ValueError):
            analyse_bond(changed, SETTLEMENT, D('5000'), D('0'))


def test_ranking_and_report_fields(bond, tmp_path):
    base = analyse_bond(bond, SETTLEMENT, D('5000'), D('10'))
    longer = replace(base, bond=replace(bond, cusip='DEMO00002', maturity_date=date(2028, 1, 15)), annualised_yield=D('.0509'))
    shorter = replace(base, annualised_yield=D('.0501'))
    best = replace(base, bond=replace(bond, cusip='DEMO00003'), annualised_yield=D('.0510'))
    rows = ranked_rows([longer, shorter, best])
    assert [r['cusip'] for r in rows] == ['DEMO00003', 'DEMO00001', 'DEMO00002']
    assert COLUMNS[:3] == ('name', 'months_to_maturity', 'net_annualised_return_pct')
    prefix = tmp_path / 'comparison'
    write_reports([base], [], prefix, 'test context')
    with prefix.with_suffix('.csv').open() as f:
        row = next(csv.DictReader(f))
    for key in COLUMNS:
        if hasattr(base, key):
            assert row[key] == str(getattr(base, key))
        elif hasattr(bond, key):
            assert row[key] == (str(getattr(bond, key)) if getattr(bond, key) is not None else '')
    assert D(row['net_annualised_return_pct']) == base.annualised_yield * 100
    sheet = load_workbook(prefix.with_suffix('.xlsx')).active
    assert sheet.max_column == len(COLUMNS)
    assert 'USD corporate' in prefix.with_suffix('.md').read_text()
    assert 'CGT' not in prefix.with_suffix('.md').read_text()
    injection = replace(base, bond=replace(bond, name='=HYPERLINK("bad")'))
    write_reports([injection], [], prefix, 'test')
    assert load_workbook(prefix.with_suffix('.xlsx')).active['A2'].data_type == 's'


def test_csv_validation(tmp_path):
    assert read_bonds(EXAMPLE)[0].moodys == 'A2'
    for content in [EXAMPLE.read_text().splitlines()[0] + '\n', 'cusip,name\nX,Y\n',
                    EXAMPLE.read_text() + EXAMPLE.read_text().splitlines()[1] + '\n']:
        path = tmp_path / 'bad.csv'
        path.write_text(content)
        with pytest.raises(ValueError):
            read_bonds(path)


def test_cli_end_to_end_and_no_matches(tmp_path):
    prefix = tmp_path / 'report'
    command = [sys.executable, '-m', 'us_corporate_analyser.cli', str(EXAMPLE),
               '--settlement-date', '2026-10-19', '--investment', '5000', '--output', str(prefix)]
    success = subprocess.run(command, capture_output=True, text=True)
    assert success.returncode == 0, success.stderr
    assert '1 ranked; 0 excluded' in success.stdout
    assert prefix.with_name('report_input.csv').read_bytes() == EXAMPLE.read_bytes()
    assert load_workbook(prefix.with_suffix('.xlsx')).active.max_row == 2
    empty = subprocess.run(command + ['--minimum-sp', 'AA-'], capture_output=True, text=True)
    assert empty.returncode == 0, empty.stderr
    assert '0 ranked; 1 excluded' in empty.stdout
    assert 'below minimum' in prefix.with_name('report_excluded.csv').read_text()
    assert len(prefix.with_suffix('.csv').read_text().splitlines()) == 1
    assert load_workbook(prefix.with_suffix('.xlsx')).active.max_row == 1
    assert subprocess.run(command + ['--months', '25'], capture_output=True).returncode != 0
    for flag in ['--investment', '--fee']:
        assert subprocess.run(command + [flag, 'Infinity', '--minimum-sp', 'A-'], capture_output=True).returncode != 0


def test_workflow_defaults_to_public_fetch_and_preserves_supplied_mode():
    text = (ROOT / '.github/workflows/run-us-corporate-analysis.yml').read_text()
    assert 'workflow_dispatch:' in text
    assert 'schedule:' not in text
    assert 'default: A2' in text
    assert 'default: A-' in text
    assert 'pytest' in text
    assert 'GITHUB_STEP_SUMMARY' in text
    assert 'actions/upload-artifact@v6' in text
    assert 'examples/' not in text
    assert 'default: Fetch public indicative research' in text
    assert 'us-corporate-fetch-public' in text
    assert 'Use repository CSV' in text


def test_summary_top_fifty_and_precision_tie(bond, tmp_path):
    base = analyse_bond(bond, SETTLEMENT, D('5000'), D('0'))
    analyses = [replace(base, bond=replace(bond, cusip=f'DEMO{i:05d}')) for i in range(51)]
    rows = ranked_rows(list(reversed(analyses)))
    assert rows[0]['cusip'] == 'DEMO00000'
    assert rows[-1]['cusip'] == 'DEMO00050'
    precise = replace(base, annualised_yield=D('.0509'))
    lower = replace(base, bond=replace(bond, cusip='AAAA00001'), annualised_yield=D('.0501'))
    assert ranked_rows([lower, precise])[0]['cusip'] == bond.cusip
    prefix = tmp_path / 'many'
    write_reports(analyses, [], prefix, 'test')
    summary = prefix.with_suffix('.md').read_text()
    assert 'DEMO00049' in summary and 'DEMO00050' not in summary
    with prefix.with_suffix('.csv').open() as handle:
        assert len(list(csv.DictReader(handle))) == 51


def test_stricter_default_rating_floor(bond):
    from us_corporate_analyser.screening import exclusion_reason
    assert exclusion_reason(bond, SETTLEMENT) == ''
    assert 'below minimum' in exclusion_reason(replace(bond, moodys='A3'), SETTLEMENT)
    assert 'below minimum' in exclusion_reason(replace(bond, sp='BBB+'), SETTLEMENT)
    assert exclusion_reason(replace(bond, moodys='Baa3', sp='BBB-'), SETTLEMENT,
                            minimum_moodys='Baa3', minimum_sp='BBB-') == ''


def test_import_source_fields_survive_reports(bond, tmp_path):
    security = replace(bond, broker_ytm_pct=D('6.5'), broker_ytw_pct=D('6.4'),
                       source_selected_nominal=D('10000'), source_accrued_interest=D('6.67'),
                       source_estimated_total=D('9906.67'))
    analysis = analyse_bond(security, SETTLEMENT, D('5000'), D('0'))
    prefix = tmp_path / 'provenance'
    write_reports([analysis], [], prefix, 'synthetic test only')
    with prefix.with_suffix('.csv').open() as handle:
        row = next(csv.DictReader(handle))
    for key in ['broker_ytm_pct', 'broker_ytw_pct', 'source_selected_nominal',
                'source_accrued_interest', 'source_estimated_total']:
        assert D(row[key]) == getattr(security, key)
        with pytest.raises(ValueError):
            validate(replace(security, **{key: D('NaN')}))
