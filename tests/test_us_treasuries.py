import csv
from dataclasses import replace
from datetime import date
from decimal import Decimal as D
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from us_treasury_analyser.calculator import Treasury, accrued_interest, analyse_treasury, coupon_schedule
from us_treasury_analyser.reporting import COLUMNS, ranked_rows, read_treasuries, write_reports


@pytest.fixture
def note():
    return Treasury("TEST-NOTE", "Test note", "note", D("4"), date(2027, 7, 15), D("98"))


def test_regular_schedule_and_coupon_day_accrual(note):
    assert coupon_schedule(note, date(2026, 7, 15)) == (date(2027, 1, 15), date(2027, 7, 15))
    assert accrued_interest(note, date(2026, 7, 15)) == 0


def test_month_end_coupon_schedule_including_leap_year(note):
    security = replace(note, maturity_date=date(2028, 2, 29))
    assert coupon_schedule(security, date(2026, 8, 31)) == (
        date(2027, 2, 28), date(2027, 8, 31), date(2028, 2, 29))
    assert accrued_interest(security, date(2027, 2, 28)) == 0


def test_mid_period_actual_actual_and_profit_separation(note):
    settlement = date(2026, 10, 15)
    result = analyse_treasury(note, settlement, D("5000"), D("10"))
    expected = D("2") * D("92") / D("184")
    assert result.accrued_interest_per_100 == expected
    assert result.dirty_price_per_100 == D("98") + expected
    assert result.nominal_purchased == D("5000")
    assert result.clean_consideration == D("4900")
    assert result.accrued_interest_paid == D("50")
    assert result.total_initial_cost == D("4960")
    assert result.unused_cash == D("40")
    assert result.coupon_income == D("200")
    assert result.redemption_gain == D("100")
    assert result.total_proceeds == D("5200")
    assert result.total_cash_profit == D("240")
    assert [flow.total for flow in result.cashflows] == [D("100"), D("5100")]


def test_treasury_regulations_reopening_example_accrual(note):
    # 31 CFR 356 Appendix B: 6.75% note, May 15 2000 maturity anchor,
    # reopened August 15 2000; May 15-November 15 is a 184-day period.
    # Compute the unrounded Actual/Actual value per $100 (broker rounding excluded).
    security = replace(note, coupon_rate=D("6.75"), maturity_date=date(2005, 5, 15))
    assert accrued_interest(security, date(2000, 8, 15)) == D("1.6875")


def test_affordability_rounds_to_broker_lots_and_includes_fee(note):
    security = replace(note, clean_buy_price=D("100"))
    result = analyse_treasury(security, date(2026, 7, 15), D("5000"), D("1"))
    assert result.nominal_purchased == D("4000")
    assert result.total_initial_cost == D("4001")
    assert result.unused_cash == D("999")
    assert result.total_initial_cost + security.nominal_increment > result.investment_budget


def test_minimum_plus_increment_grid(note):
    security = replace(note, clean_buy_price=D("100"), minimum_nominal=D("2000"), nominal_increment=D("1000"))
    assert analyse_treasury(security, date(2026, 7, 15), D("3500"), D("0")).nominal_purchased == D("3000")
    with pytest.raises(ValueError, match="minimum_nominal"):
        analyse_treasury(security, date(2026, 7, 15), D("1999"), D("0"))


def test_bill_has_single_redemption_and_actual365_yield(note):
    bill = replace(note, security_type="bill", coupon_rate=D("0"), maturity_date=date(2027, 1, 1), clean_buy_price=D("95"))
    result = analyse_treasury(bill, date(2026, 1, 1), D("1000"), D("0"))
    assert result.accrued_interest_per_100 == 0
    assert result.coupon_income == 0
    assert len(result.cashflows) == 1
    assert result.cashflows[0].redemption == D("1000")
    assert result.total_initial_cost == D("950")
    assert result.annualised_yield == pytest.approx(D("1000") / D("950") - 1, abs=D("1e-12"))
    fee_result = analyse_treasury(bill, date(2026, 1, 1), D("1000"), D("10"))
    assert fee_result.annualised_yield < result.annualised_yield


def test_coupon_bond_yield_solves_dated_flows(note):
    security = replace(note, security_type="bond", clean_buy_price=D("100"), maturity_date=date(2027, 1, 1))
    result = analyse_treasury(security, date(2026, 1, 1), D("1000"), D("0"))
    present_value = sum((flow.total / ((1 + result.annualised_yield) ** (D((flow.payment_date - result.settlement_date).days) / D("365"))) for flow in result.cashflows), D("0"))
    assert present_value == pytest.approx(result.total_initial_cost, abs=D("1e-10"))
    assert D("0.04") < result.annualised_yield < D("0.041")


@pytest.mark.parametrize("changes", [
    {"clean_buy_price": D("0")}, {"clean_buy_price": D("NaN")},
    {"coupon_rate": D("Infinity")}, {"coupon_rate": D("-1")},
    {"coupon_rate": 4.0}, {"minimum_nominal": D("0")},
    {"nominal_increment": D("-1")}, {"security_type": "tips"},
    {"security_type": "frn"}, {"security_type": "strips"},
    {"security_type": "bill"}, {"cusip": ""}, {"name": ""},
    {"coupon_rate": D("0")}, {"quote_date": "bad-date"},
])
def test_invalid_security_inputs(note, changes):
    with pytest.raises(ValueError):
        analyse_treasury(replace(note, **changes), date(2026, 7, 15), D("5000"), D("0"))


@pytest.mark.parametrize("budget,fee", [(D("0"), D("0")), (D("NaN"), D("0")), (D("5000"), D("NaN")), (D("5000"), D("-1")), (D("5000"), D("5000"))])
def test_invalid_budget_and_fee(note, budget, fee):
    with pytest.raises(ValueError):
        analyse_treasury(note, date(2026, 7, 15), budget, fee)


def test_matured_security_rejected(note):
    with pytest.raises(ValueError, match="before maturity"):
        analyse_treasury(note, note.maturity_date, D("5000"), D("0"))


def test_report_ranking_bands_shorter_first_and_negative_yields(note):
    item = analyse_treasury(note, date(2026, 7, 15), D("5000"), D("0"))
    shorter = replace(item, treasury=replace(note, cusip="SHORT", maturity_date=date(2027, 1, 15)), annualised_yield=D("0.0421"))
    longer = replace(item, annualised_yield=D("0.0429"))
    higher_band = replace(item, treasury=replace(note, cusip="HIGH"), annualised_yield=D("0.043"))
    negative = replace(item, treasury=replace(note, cusip="NEG"), annualised_yield=D("-0.0001"))
    assert [row["cusip"] for row in ranked_rows([longer, negative, shorter, higher_band])] == ["HIGH", "SHORT", "TEST-NOTE", "NEG"]


def test_csv_excel_fields_and_decimal_roundtrip(note, tmp_path):
    item = analyse_treasury(note, date(2026, 10, 15), D("5000"), D("10"))
    csv_path, xlsx_path = write_reports([item], tmp_path / "us_treasuries")
    with csv_path.open() as handle:
        reader = csv.DictReader(handle)
        row = next(reader)
        assert tuple(reader.fieldnames) == COLUMNS
        assert tuple(reader.fieldnames[:6]) == ("cusip", "security_id", "name", "months_to_maturity", "net_annualised_return_pct", "rank")
    assert row["currency"] == "USD"
    assert D(row["total_cash_profit"]) == item.total_cash_profit
    assert D(row["net_annualised_return_pct"]) == item.annualised_yield * 100
    assert "cgt_exempt_redemption_gain" not in row
    workbook = load_workbook(xlsx_path)
    assert tuple(cell.value for cell in workbook.active[1]) == COLUMNS
    assert workbook.active.max_row == 2


@pytest.mark.parametrize("content", ["cusip,name\n", "cusip,name,security_type,coupon_rate,maturity_date,clean_buy_price\n", "cusip,name,security_type,coupon_rate,maturity_date,clean_buy_price\nA,Test,note,NaN,2027-01-01,100\n", "cusip,name,security_type,coupon_rate,maturity_date,clean_buy_price\nA,Test,note,4,2027-01-01\n"])
def test_invalid_csv(tmp_path, content):
    path = tmp_path / "bad.csv"
    path.write_text(content)
    with pytest.raises(ValueError):
        read_treasuries(path)


def test_example_and_cli_end_to_end(tmp_path):
    example = Path(__file__).resolve().parents[1] / "examples/us_treasuries.csv"
    assert len(read_treasuries(example)) == 3
    result = subprocess.run([sys.executable, "-m", "us_treasury_analyser.cli", str(example), "--investment", "5000", "--settlement-date", "2026-10-14", "--output", str(tmp_path / "us_treasuries")], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert (tmp_path / "us_treasuries.csv").is_file()
    assert (tmp_path / "us_treasuries.xlsx").is_file()


def test_maximum_caps_purchase_on_minimum_increment_grid(note):
    security = replace(note, clean_buy_price=D("100"), minimum_nominal=D("2000"), maximum_nominal=D("4500"))
    result = analyse_treasury(security, date(2026, 7, 15), D("10000"), D("0"))
    assert result.nominal_purchased == D("4000")
    assert result.unused_cash == D("6000")


@pytest.mark.parametrize("changes", [
    {"maximum_nominal": D("999")}, {"maximum_nominal": D("NaN")},
    {"maximum_nominal": D("Infinity")}, {"maximum_nominal": 2000.0},
    {"broker_ytm_pct": D("NaN")}, {"broker_ytm_pct": D("Infinity")},
    {"broker_ytm_pct": 4.2},
])
def test_invalid_quote_limits_and_ytm(note, changes):
    with pytest.raises(ValueError):
        analyse_treasury(replace(note, **changes), date(2026, 7, 15), D("5000"), D("0"))


def test_quote_limits_and_broker_ytm_report_fields(note, tmp_path):
    security = replace(note, maximum_nominal=D("10000"), broker_ytm_pct=D("4.207"))
    result = analyse_treasury(security, date(2026, 7, 15), D("5000"), D("0"))
    csv_path, xlsx_path = write_reports([result], tmp_path / "quotes")
    with csv_path.open() as handle:
        row = next(csv.DictReader(handle))
    assert row["maximum_nominal"] == "10000"
    assert row["broker_ytm_pct"] == "4.207"
    assert D(row["net_annualised_return_pct"]) != D(row["broker_ytm_pct"])
    values = list(load_workbook(xlsx_path).active.values)
    assert values[1][COLUMNS.index("broker_ytm_pct")] == 4.207


def test_schwab_snapshot_converts_units_and_preserves_all_conventional_quotes():
    path = Path(__file__).resolve().parents[1] / "examples/us_treasuries_schwab.csv"
    securities = read_treasuries(path)
    assert len(securities) == 14
    indexed = {item.cusip: item for item in securities}
    assert indexed["912797WR4"].minimum_nominal == D("1000")
    assert indexed["912797WR4"].maximum_nominal == D("15001000")
    assert indexed["912797WR4"].clean_buy_price == D("98.8267")
    assert indexed["912797WR4"].broker_ytm_pct == D("4.207")
    assert indexed["912797WQ6"].minimum_nominal == D("2700000")
    assert indexed["91282CJT9"].minimum_nominal == D("200000")
    assert indexed["91282CME8"].coupon_rate == D("4.250")
    assert "912828V49" not in indexed and "912810PS1" not in indexed


def test_schwab_cli_reports_four_eligible_and_ten_unaffordable(tmp_path):
    path = Path(__file__).resolve().parents[1] / "examples/us_treasuries_schwab.csv"
    result = subprocess.run([sys.executable, "-m", "us_treasury_analyser.cli", str(path), "--investment", "5000", "--settlement-date", "2026-10-15", "--skip-unaffordable", "--output", str(tmp_path / "us_treasuries")], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    with (tmp_path / "us_treasuries.csv").open() as handle:
        eligible = list(csv.DictReader(handle))
    with (tmp_path / "us_treasuries_excluded.csv").open() as handle:
        excluded = list(csv.DictReader(handle))
    assert len(eligible) == 4
    assert len(excluded) == 10
    assert {row["cusip"] for row in eligible} == {"912797WR4", "912797TM9", "912797VU8", "912797VS3"}


@pytest.mark.parametrize("budget,skip,expected", [("5000", False, 2), ("500", True, 0), ("NaN", True, 2)])
def test_cli_exclusions_do_not_hide_invalid_inputs(tmp_path, budget, skip, expected):
    path = Path(__file__).resolve().parents[1] / "examples/us_treasuries_schwab.csv"
    command = [sys.executable, "-m", "us_treasury_analyser.cli", str(path), "--investment", budget, "--settlement-date", "2026-10-15", "--output", str(tmp_path / "result")]
    if skip:
        command.append("--skip-unaffordable")
    result = subprocess.run(command, capture_output=True, text=True)
    assert result.returncode == expected, result.stderr
    if expected == 0:
        with (tmp_path / "result.csv").open() as handle:
            assert list(csv.DictReader(handle)) == []
        with (tmp_path / "result_excluded.csv").open() as handle:
            assert len(list(csv.DictReader(handle))) == 14
    else:
        assert not (tmp_path / "result.csv").exists()
