import csv
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from gilt_analyser.calculator import analyse_gilt
from gilt_analyser.models import Gilt
from gilt_analyser.reporting import ranked_rows, read_gilts, write_reports


def test_read_and_write_ranked_reports(tmp_path) -> None:
    source = tmp_path / "input.csv"
    source.write_text(
        "isin,hl_code,name,coupon_rate,maturity_date,clean_buy_price\n"
        "LOW,LOW27,Low return,0,2027-01-01,100\n"
        "HIGH,HIGH27,High return,0,2027-01-01,95\n",
        encoding="utf-8",
    )
    analyses = [analyse_gilt(g, date(2026, 1, 1), Decimal("1000"), Decimal("0")) for g in read_gilts(source)]
    csv_path, xlsx_path = write_reports(analyses, tmp_path / "comparison")
    with csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert list(rows[0])[3:6] == ["name", "months_to_maturity", "net_annualised_return_pct"]
    assert [row["isin"] for row in rows] == ["HIGH", "LOW"]
    assert rows[0]["hl_code"] == "HIGH27"
    assert rows[0]["cgt_exempt_redemption_gain"] != ""
    assert rows[0]["months_to_maturity"] == "12.00"
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook["Gilt comparison"]["B2"].value == "HIGH"
    assert workbook["Gilt comparison"]["C2"].value == "HIGH27"
    sheet = workbook["Gilt comparison"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(rows[0])
    assert sheet["F2"].value == pytest.approx(float(rows[0]["net_annualised_return_pct"]))
    assert sheet.cell(2, headers.index("months_to_maturity") + 1).value == 12


@pytest.mark.parametrize(
    "settlement,maturity,expected",
    [
        (date(2026, 1, 1), date(2026, 1, 2), Decimal("0.03")),
        (date(2026, 1, 1), date(2026, 7, 1), Decimal("5.95")),
        (date(2027, 3, 1), date(2028, 3, 1), Decimal("12.03")),
    ],
)
def test_months_to_maturity_uses_actual_days(settlement, maturity, expected) -> None:
    analysis = analyse_gilt(
        Gilt("TEST", "Test", Decimal("0"), maturity, Decimal("100")),
        settlement, Decimal("1000"), Decimal("0"),
    )
    assert ranked_rows([analysis])[0]["months_to_maturity"] == expected


def test_same_total_return_ranks_shorter_maturity_first() -> None:
    analyses = [
        analyse_gilt(
            Gilt(isin, isin, Decimal("0"), maturity, Decimal("95")),
            date(2026, 1, 1), Decimal("1000"), Decimal("0"),
        )
        for isin, maturity in [
            ("LONG", date(2028, 1, 1)), ("SHORT", date(2027, 1, 1)),
        ]
    ]
    assert analyses[0].total_cash_profit == analyses[1].total_cash_profit
    assert [row["isin"] for row in ranked_rows(analyses)] == ["SHORT", "LONG"]


@pytest.mark.parametrize(
    "short_yield,long_yield,expected",
    [
        ("0.0420", "0.0429", ["SHORT", "LONG"]),
        ("0.0429", "0.0430", ["LONG", "SHORT"]),
        ("0.0420", "0.0420", ["SHORT", "LONG"]),
        ("-0.0009", "-0.0001", ["SHORT", "LONG"]),
        ("-0.0001", "0", ["LONG", "SHORT"]),
    ],
)
def test_similar_yield_bands_prefer_shorter_maturity(short_yield, long_yield, expected) -> None:
    base = analyse_gilt(
        Gilt("LONG", "Long", Decimal("0"), date(2028, 1, 1), Decimal("95")),
        date(2026, 1, 1), Decimal("1000"), Decimal("0"),
    )
    longer = replace(base, annualised_yield=Decimal(long_yield))
    shorter = replace(
        base, gilt=replace(base.gilt, isin="SHORT", maturity_date=date(2027, 1, 1)),
        annualised_yield=Decimal(short_yield),
    )
    for inputs in ([longer, shorter], [shorter, longer]):
        rows = ranked_rows(inputs)
        assert [row["isin"] for row in rows] == expected
        assert [row["rank"] for row in rows] == [1, 2]


def test_input_without_optional_hl_code_remains_supported(tmp_path) -> None:
    source = tmp_path / "input.csv"
    source.write_text(
        "isin,name,coupon_rate,maturity_date,clean_buy_price\n"
        "OLD,Older input,0,2027-01-01,99\n",
        encoding="utf-8",
    )
    assert read_gilts(source)[0].hl_code == ""
