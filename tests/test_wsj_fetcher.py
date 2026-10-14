import csv
import json
import sys
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
from openpyxl import load_workbook

from us_treasury_analyser import cli, wsj_fetcher
from us_treasury_analyser.calculator import analyse_treasury
from us_treasury_analyser.reporting import COLUMNS, read_treasuries, write_reports


@pytest.fixture
def tables():
    return {
        "BILLS": {"timestamp": "Tuesday, October 13, 2026", "instruments": [
            {"maturityDate": "4/15/2027", "ask": "4.185", "askYield": "4.335", "bid": "4.195"},
        ]},
        "NOTES_AND_BONDS": {"timestamp": "Wednesday, October 14, 2026", "instruments": [
            {"maturityDate": "4/30/2027", "coupon": "0.500", "ask": "97.3120", "askYield": "4.306"},
            {"maturityDate": "4/30/2027", "coupon": "2.750", "ask": "99.0520", "askYield": "4.321"},
            {"maturityDate": "4/15/2027", "coupon": "4.500", "ask": "100.0260", "askYield": "4.324"},
        ]},
    }


def page(tables):
    data = {"mdc_treasury_" + kind: {"data": {"id": json.dumps({"treasury": kind}), "data": table}}
            for kind, table in tables.items()}
    return "<script>window.__STATE__ = " + json.dumps({"data": data}) + ";</script>"


def test_embedded_state_extraction(tables):
    assert wsj_fetcher.extract_tables(page(tables)) == tables


@pytest.mark.parametrize("html", ["blocked", "window.__STATE__ = {}", "window.__STATE__ = []", "window.__STATE__ = invalid"])
def test_invalid_page_rejected(html):
    with pytest.raises(ValueError):
        wsj_fetcher.extract_tables(html)


def test_missing_and_empty_tables_rejected(tables):
    del tables["BILLS"]
    with pytest.raises(ValueError, match="missing"):
        wsj_fetcher.extract_tables(page(tables))


@pytest.mark.parametrize("quote,price", [
    ("99.0320", "99.1015625"),  # CME cash-market 99-032 worked example.
    ("99.0340", "99.109375"),  # CME 99-03+ example (4/8 of a 32nd).
    ("100.0260", "100.0859375"),
    ("97.3120", "97.9765625"),
    ("99.0520", "99.1640625"),
    ("100.0000", "100"),
    ("88.2370", "88.74609375"),
])
def test_fractional_price_conversion(quote, price):
    assert wsj_fetcher.treasury_fractional_price(quote) == D(price)


@pytest.mark.parametrize("quote", ["99.32", "99.3200", "99.0380", "99.0321", "99-032", "NaN", "n.a.", "0.0000", 99.032])
def test_fractional_price_rejects_ambiguous_or_invalid_values(quote):
    with pytest.raises(ValueError):
        wsj_fetcher.treasury_fractional_price(quote)


def test_bill_conversion_matches_treasury_direct_example():
    # TreasuryDirect's $1,000 / 182-day / 0.145% example rounds to $999.27.
    result = wsj_fetcher.bill_price(D("0.145"), date(2026, 1, 1), date(2026, 7, 2))
    assert (result * 10).quantize(D("0.01")) == D("999.27")
    assert result == D("100") - D("0.145") * D("182") / D("360")


def test_bill_conversion_matches_treasury_regulations_example():
    # Official 90-day, 7.61% example: unrounded $98.0975 per $100.
    assert wsj_fetcher.bill_price(D("7.61"), date(1989, 11, 24), date(1990, 2, 22)) == D("98.0975")


@pytest.mark.parametrize("rate", [D("NaN"), D("Infinity"), D("10000"), 4.185])
def test_bad_bill_rates_rejected(rate):
    with pytest.raises(ValueError):
        wsj_fetcher.bill_price(rate, date(2026, 10, 15), date(2027, 4, 15))


def test_bill_conversion_requires_future_maturity():
    with pytest.raises(ValueError, match="before maturity"):
        wsj_fetcher.bill_price(D("4"), date(2027, 4, 15), date(2027, 4, 15))


def test_terms_and_source_metadata_preserved_without_inventing_cusips(tables):
    securities = wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15), fetched_at_utc="2026-10-14T16:00:00+00:00")
    assert len(securities) == 4
    bill = securities[0]
    assert bill.cusip == ""
    assert bill.security_id == "WSJ:bill:2027-04-15:0"
    assert bill.quote_date == "2026-10-13"
    assert bill.source_ask_yield_pct == D("4.335")
    assert bill.broker_ytm_pct is None
    assert bill.quote_status == "indicative" and bill.quantity_basis == "assumed"
    assert bill.maximum_nominal is None
    assert bill.price_settlement_date == "2026-10-15"
    assert bill.clean_buy_price == D("100") - D("4.185") * D("182") / D("360")
    assert securities[1].clean_buy_price == D("97.9765625")
    assert securities[1].security_type == "note_or_bond"


def test_twenty_four_month_window_boundaries(tables):
    base = tables["BILLS"]["instruments"][0]
    tables["BILLS"]["instruments"] = [dict(base, maturityDate=value) for value in (
        "10/15/2026", "10/16/2026", "10/15/2028", "10/16/2028")]
    securities = wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))
    assert [s.maturity_date for s in securities if s.security_type == "bill"] == [date(2026, 10, 16), date(2028, 10, 15)]


def test_calendar_month_window_clamps_leap_and_nonleap_year():
    assert wsj_fetcher._add_months(date(2027, 1, 31), 1) == date(2027, 2, 28)
    assert wsj_fetcher._add_months(date(2028, 1, 31), 1) == date(2028, 2, 29)


@pytest.mark.parametrize("settlement,age", [(date(2026, 10, 22), 7), (date(2026, 10, 12), 7), (date(2026, 10, 15), 0)])
def test_stale_future_or_over_age_quotes_rejected(tables, settlement, age):
    with pytest.raises(ValueError, match="quote date"):
        wsj_fetcher.parse_treasuries(tables, settlement, max_quote_age_days=age)


@pytest.mark.parametrize("changes", [{"ask": "n.a."}, {"askYield": "NaN"}, {"coupon": "-1"}, {"coupon": "0"}, {"maturityDate": "invalid"}])
def test_malformed_or_unsupported_coupon_quotes_fail(tables, changes):
    tables["NOTES_AND_BONDS"]["instruments"][0].update(changes)
    with pytest.raises(ValueError):
        wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))


def test_duplicate_terms_fail_instead_of_fabricating_distinct_securities(tables):
    tables["BILLS"]["instruments"] *= 2
    with pytest.raises(ValueError, match="ambiguous duplicate"):
        wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))


def test_normalized_input_roundtrip_and_report_fields(tables, tmp_path):
    securities = wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))
    input_path = tmp_path / "input.csv"
    wsj_fetcher.write_input_csv(securities, input_path)
    assert read_treasuries(input_path) == securities
    analyses = [analyse_treasury(s, date(2026, 10, 15), D("5000"), D("0")) for s in securities]
    csv_path, xlsx_path = write_reports(analyses, tmp_path / "comparison")
    with csv_path.open() as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 4
    for row in rows:
        assert row["cusip"] == row["minimum_nominal"] == row["nominal_increment"] == row["maximum_nominal"] == ""
        assert row["analysis_minimum_nominal"] == row["analysis_nominal_increment"] == "1000"
        assert row["quantity_basis"] == "assumed"
        assert row["quote_status"] == "indicative"
        assert row["source_ask_quote"]
        assert row["source_ask_yield_pct"]
        assert row["price_basis"]
        assert row["broker_ytm_pct"] == ""
    sheet = load_workbook(xlsx_path).active
    assert sheet.max_row == 5
    assert [cell.value for cell in sheet[1]] == list(COLUMNS)


def test_derived_bill_price_cannot_be_reused_for_different_settlement(tables):
    bill = wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))[0]
    with pytest.raises(ValueError, match="fetch quotes again"):
        analyse_treasury(bill, date(2026, 10, 16), D("5000"), D("0"))


def test_fetch_reads_public_page_with_tls_and_saves_snapshot(tables, monkeypatch, tmp_path):
    class Response:
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def read(self):
            return page(tables).encode()
    def open_page(request, timeout, context):
        assert request.full_url == wsj_fetcher.WSJ_URL
        assert context.check_hostname
        assert timeout == 30
        return Response()
    monkeypatch.setattr(wsj_fetcher, "urlopen", open_page)
    snapshot = tmp_path / "source.json"
    securities = wsj_fetcher.fetch_treasuries(date(2026, 10, 15), snapshot_path=snapshot)
    assert len(securities) == 4
    assert json.loads(snapshot.read_text())["tables"] == tables


def test_one_command_fetch_to_csv_excel_and_input_snapshot(tables, monkeypatch, tmp_path):
    securities = wsj_fetcher.parse_treasuries(tables, date(2026, 10, 15))
    def fetch(settlement, months, max_age, snapshot):
        assert settlement == date(2026, 10, 15) and months == 24 and max_age == 7
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        snapshot.write_text(json.dumps(tables))
        return securities
    monkeypatch.setattr(cli, "fetch_treasuries", fetch)
    output = tmp_path / "reports/us_treasuries"
    monkeypatch.setattr(sys, "argv", ["us-treasury-analyser", "--investment", "5000", "--settlement-date", "2026-10-15", "--output", str(output)])
    cli.main()
    assert output.with_suffix(".csv").exists() and output.with_suffix(".xlsx").exists()
    assert read_treasuries(output.with_name("us_treasuries_input.csv")) == securities
    assert output.with_name("us_treasuries_source.json").exists()


@pytest.mark.parametrize("args", [["--months", "0"], ["--max-quote-age-days", "-1"], ["example.csv", "--fetch-wsj"]])
def test_bad_fetch_options_fail_before_network(monkeypatch, args):
    monkeypatch.setattr(sys, "argv", ["us-treasury-analyser", "--investment", "5000", "--settlement-date", "2026-10-15", *args])
    with pytest.raises(SystemExit) as error:
        cli.main()
    assert error.value.code == 2
